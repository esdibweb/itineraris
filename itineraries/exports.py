"""Excel export of students, their hours and the activities."""
import datetime
from io import BytesIO

from django.db.models import Count, Sum
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gcodex.models import Student
from .models import Activity, Annotation, StudentHours

DATE_FORMAT = 'DD/MM/YYYY'
MAX_COLUMN_WIDTH = 60

# The concrete activity type is a one-to-one child of Activity
ACTIVITY_TYPES = [
    ('activityinternalcolaborate', 'Col·laboració interna'),
    ('activityaward', 'Premi'),
    ('activityinternalrepresent', 'Representació'),
    ('activityexternal', 'Col·laboració externa'),
]
# Related rows needed to show each activity's type and detail without extra queries
ACTIVITY_RELATIONS = [relation for relation, _ in ACTIVITY_TYPES] + [
    'activityinternalcolaborate__teacher', 'activityinternalrepresent__type',
]


def activity_type_and_detail(activity):
    """(type label, organisation / teacher / representation type) of an activity."""
    for relation, label in ACTIVITY_TYPES:
        child = getattr(activity, relation, None)
        if child is None:
            continue
        if relation == 'activityinternalcolaborate':
            detail = child.teacher.complete_name if child.teacher else ''
        elif relation == 'activityinternalrepresent':
            detail = child.type.description if child.type else ''
        else:
            detail = child.organitzation
        return label, detail
    return '', ''


def add_sheet(workbook, title, headers, rows):
    """Add a sheet with a bold, frozen, filterable header row and sized columns."""
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    widths = [len(header) for header in headers]
    for row in rows:
        sheet.append(row)
        for index, value in enumerate(row):
            widths[index] = max(widths[index], len(str(value)) if value is not None else 0)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            # Text such as "=HYPERLINK(...)" must stay text, never become a formula
            if cell.data_type == 'f':
                cell.data_type = 's'
            elif isinstance(cell.value, datetime.date):
                cell.number_format = DATE_FORMAT

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def students_rows():
    totals = dict(
        StudentHours.objects.order_by().values('student').annotate(total=Sum('hours')).values_list('student', 'total')
    )
    annotations = {annotation.student_id: annotation for annotation in Annotation.objects.all()}
    students = Student.objects.filter(pk__in=set(totals) | set(annotations)).order_by('complete_name')
    for student in students:
        annotation = annotations.get(student.pk)
        yield [
            student.id_number,
            student.complete_name,
            student.school_email,
            totals.get(student.pk, 0),
            annotation.date if annotation else None,
            annotation.observations if annotation else None,
        ]


def hours_rows():
    student_hours = StudentHours.objects.select_related(
        'student', *(f'activity__{relation}' for relation in ACTIVITY_RELATIONS),
    ).order_by('student__complete_name', 'activity__date')
    for hours in student_hours:
        activity_type, detail = activity_type_and_detail(hours.activity)
        yield [
            hours.student.id_number,
            hours.student.complete_name,
            hours.activity.title,
            activity_type,
            hours.activity.date,
            detail,
            hours.hours,
        ]


def activities_rows():
    activities = Activity.objects.select_related(*ACTIVITY_RELATIONS).annotate(
        student_count=Count('hours__student', distinct=True), total_hours=Sum('hours__hours'),
    ).order_by('-date', 'title')
    for activity in activities:
        activity_type, detail = activity_type_and_detail(activity)
        yield [
            activity.title,
            activity_type,
            activity.date,
            detail,
            activity.student_count,
            activity.total_hours or 0,
        ]


def build_workbook():
    """Return the export as .xlsx bytes."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(
        workbook, 'Alumnat',
        ['DNI', 'Nom', "Correu de l'escola", "Total d'hores", 'Data pagament crèdits', 'Observacions'],
        list(students_rows()),
    )
    add_sheet(
        workbook, 'Hores',
        ['DNI', 'Nom', 'Activitat', 'Tipus', 'Data', 'Entitat / docent / tipus', 'Hores'],
        list(hours_rows()),
    )
    add_sheet(
        workbook, 'Activitats',
        ['Activitat', 'Tipus', 'Data', 'Entitat / docent / tipus', "Nombre d'alumnes", "Total d'hores"],
        list(activities_rows()),
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
